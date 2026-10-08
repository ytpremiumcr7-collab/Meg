"""Saved TIN -> volume persistence, using the real domain service and database."""
from uuid import uuid4

import pytest
from sqlalchemy import bindparam, func, select, text
from sqlalchemy.exc import IntegrityError
from app.db.types import UUID as DBUUID

from app.core.errors import MegalodonException
from app.models.expediente import ExpedienteObra
from app.models.topografia import CalculoVolumen, Levantamiento, SuperficieTIN
from app.services.topografia_service import TopografiaService
from app.services.presupuesto_service import PresupuestoService
from app.models.presupuesto import Presupuesto, Partida
from tests.integration.test_bim_regressions import PARAMS
from tests.conftest import AsyncSessionLocalTest


@pytest.mark.asyncio
@pytest.mark.parametrize("crs,srid", [("EPSG:4326", 4326), ("EPSG:2277", 2277), ("EPSG:6362", 4326)])
async def test_volume_rejects_nonmetric_or_inconsistent_crs(db_session, tenant_a_user, crs, srid):
    tenant, user = tenant_a_user
    existing, _ = await surfaces(db_session, tenant, user)
    survey = await db_session.get(Levantamiento, existing.levantamiento_id)
    survey.crs, survey.srid = crs, srid
    await db_session.commit()
    with pytest.raises(MegalodonException, match="CRS|métric|metros"):
        await TopografiaService(db_session, tenant.id).calcular_volumen(
            superficie_existente_id=existing.id, elevacion_referencia=0,
        )


@pytest.mark.asyncio
async def test_volume_rejects_two_different_coordinate_frames(db_session, tenant_a_user):
    tenant, user = tenant_a_user
    existing, project = await surfaces(db_session, tenant, user)
    survey = Levantamiento(expediente_id=existing.expediente_id, identificador=uuid4().hex,
                           nombre="Other zone", crs="EPSG:32615", srid=32615)
    db_session.add(survey)
    await db_session.flush()
    project.levantamiento_id = survey.id
    await db_session.commit()
    with pytest.raises(MegalodonException, match="CRS|referencia"):
        await TopografiaService(db_session, tenant.id).calcular_volumen(
            superficie_existente_id=existing.id, superficie_proyecto_id=project.id,
        )


@pytest.mark.asyncio
async def test_reference_and_project_are_mutually_exclusive(db_session, tenant_a_user):
    tenant, user = tenant_a_user
    existing, project = await surfaces(db_session, tenant, user)
    with pytest.raises(MegalodonException, match="una|excluy|ambas"):
        await TopografiaService(db_session, tenant.id).calcular_volumen(
            superficie_existente_id=existing.id, superficie_proyecto_id=project.id, elevacion_referencia=0,
        )


@pytest.mark.asyncio
@pytest.mark.parametrize('crs,srid', [('EPSG:6362', 32614), ('EPSG:999999999', 999999999)])
async def test_invalid_coordinate_declaration_is_rejected_before_survey_creation(db_session, tenant_a_user, crs, srid):
    tenant, user = tenant_a_user
    existing, _ = await surfaces(db_session, tenant, user)
    with pytest.raises(MegalodonException, match='CRS|SRID'):
        await TopografiaService(db_session, tenant.id).crear_levantamiento(
            expediente_id=existing.expediente_id, nombre='Invalid declaration', crs=crs, srid=srid,
        )


async def surfaces(db, tenant, user):
    exp = ExpedienteObra(
        tenant_id=tenant.id, responsable_id=user.id, identificador=uuid4().hex,
        titulo="TIN regression", organo="Dependencia", unidad_administrativa="UA",
        serie_documental="OBRA", subserie_documental="LICITACION", proyecto_nombre="TIN",
    )
    db.add(exp)
    await db.flush()
    survey = Levantamiento(expediente_id=exp.id, identificador=uuid4().hex, nombre="TIN")
    db.add(survey)
    await db.flush()
    common = dict(
        expediente_id=exp.id, levantamiento_id=survey.id,
        area_plan_m2=1, area_superficie_m2=1, elevacion_min=0,
        elevacion_max=1, elevacion_media=.25, num_puntos=4, num_triangulos=2,
    )
    existing = SuperficieTIN(
        **common, nombre="Existing", malla_vertices=[0,0,1, 1,0,0, 1,1,0, 0,1,0],
        malla_caras=[2,3,1, 1,3,0],
    )
    project = SuperficieTIN(
        **common, nombre="Project", malla_vertices=[0,0,0, 1,0,0, 1,1,0, 0,1,0],
        malla_caras=[0,1,2, 0,2,3],
    )
    db.add_all([existing, project])
    await db.commit()
    return existing, project


@pytest.mark.asyncio
@pytest.mark.parametrize("against_surface", [False, True])
async def test_domain_persists_volume_of_the_saved_faces(db_session, tenant_a_user, against_surface):
    tenant, user = tenant_a_user
    existing, project = await surfaces(db_session, tenant, user)
    args = ({"superficie_proyecto_id": project.id} if against_surface else {"elevacion_referencia": 0})
    calculation = await TopografiaService(db_session, tenant.id).calcular_volumen(
        superficie_existente_id=existing.id, **args,
    )
    await db_session.refresh(calculation)
    assert float(calculation.volumen_corte_m3) == .167
    assert float(calculation.volumen_terraplen_m3) == 0
    assert float(calculation.area_analizada_m2) == 1
    assert calculation.evidencia['cobertura']['completa'] is True


@pytest.mark.asyncio
async def test_domain_does_not_persist_success_for_corrupt_faces(db_session, tenant_a_user):
    tenant, user = tenant_a_user
    existing, _ = await surfaces(db_session, tenant, user)
    existing.malla_caras = [0, 1, 9]
    await db_session.commit()
    with pytest.raises(MegalodonException):
        await TopografiaService(db_session, tenant.id).calcular_volumen(
            superficie_existente_id=existing.id, elevacion_referencia=0,
        )
    count = await db_session.scalar(select(func.count()).select_from(CalculoVolumen).where(
        CalculoVolumen.superficie_existente_id == existing.id,
    ))
    assert count == 0


@pytest.mark.asyncio
async def test_survey_with_missing_z_cannot_become_a_flat_surface(db_session, tenant_a_user):
    tenant, user = tenant_a_user
    existing, _ = await surfaces(db_session, tenant, user)
    postgres = db_session.bind.dialect.name == 'postgresql'
    geom = 'ST_GeomFromEWKT(:geom)' if postgres else ':geom'
    statement = text(f'''INSERT INTO puntos_topograficos
        (id, levantamiento_id, identificador, x, y, z, geom, precision_xy)
        VALUES (:id, :survey, :name, :x, :y, :z, {geom}, 0.02)''').bindparams(
        bindparam('id', type_=DBUUID(as_uuid=True)), bindparam('survey', type_=DBUUID(as_uuid=True)),
    )
    for i, (x, y, z) in enumerate([(0, 0, 1), (1, 0, None), (0, 1, 1)]):
        await db_session.execute(statement, dict(id=uuid4(), survey=existing.levantamiento_id,
            name=f'P{i}', x=x, y=y, z=z, geom=f'SRID=6362;POINT({x} {y})'))
    await db_session.commit()
    before = await db_session.scalar(select(func.count()).select_from(SuperficieTIN).where(
        SuperficieTIN.levantamiento_id == existing.levantamiento_id,
    ))
    with pytest.raises(MegalodonException, match="elevaci|Z") as failure:
        await TopografiaService(db_session, tenant.id).triangular_superficie(existing.levantamiento_id, 'Invalid')
    assert failure.value.details['puntos_sin_elevacion'] == ['P1']
    after = await db_session.scalar(select(func.count()).select_from(SuperficieTIN).where(
        SuperficieTIN.levantamiento_id == existing.levantamiento_id,
    ))
    assert after == before


@pytest.mark.asyncio
async def test_postgis_keeps_survey_srid_and_rejects_direct_frame_drift(db_session, tenant_a_user):
    if db_session.bind.dialect.name != 'postgresql':
        pytest.skip('Requires real PostGIS typmod and deferred constraint triggers')
    tenant, user = tenant_a_user
    existing, _ = await surfaces(db_session, tenant, user)
    svc = TopografiaService(db_session, tenant.id)
    survey = await svc.crear_levantamiento(expediente_id=existing.expediente_id, nombre='UTM',
                                           crs='EPSG:32614', srid=32614)
    points = await svc.agregar_puntos(survey.id, [dict(identificador='P1', x=500000, y=2200000, z=None)])
    point_id, survey_id = points[0].id, survey.id
    row = (await db_session.execute(text('SELECT ST_SRID(geom), ST_NDims(geom), z FROM puntos_topograficos WHERE id=:id')
        .bindparams(bindparam('id', type_=DBUUID(as_uuid=True))), dict(id=point_id))).one()
    assert tuple(row) == (32614, 2, None)
    known = await svc.agregar_puntos(survey_id, [dict(identificador='P2', x=500001, y=2200000, z=123.456)])
    known_id = known[0].id
    row = (await db_session.execute(text('SELECT ST_SRID(geom), ST_NDims(geom), ST_Z(geom), z FROM puntos_topograficos WHERE id=:id')
        .bindparams(bindparam('id', type_=DBUUID(as_uuid=True))), dict(id=known_id))).one()
    assert tuple(row[:2]) == (32614, 3)
    assert float(row[2]) == pytest.approx(123.456)
    assert float(row[3]) == pytest.approx(123.456)
    await db_session.rollback()

    with pytest.raises(IntegrityError, match='ck_punto_tipo_geometria'):
        await db_session.execute(text("UPDATE puntos_topograficos SET geom=ST_GeomFromText('LINESTRING(0 0,1 1)',32614) WHERE id=:id")
            .bindparams(bindparam('id', type_=DBUUID(as_uuid=True))), dict(id=point_id))
        await db_session.commit()
    await db_session.rollback()
    with pytest.raises(IntegrityError, match='SRID'):
        await db_session.execute(text('UPDATE puntos_topograficos SET geom=ST_SetSRID(geom,6362) WHERE id=:id')
            .bindparams(bindparam('id', type_=DBUUID(as_uuid=True))), dict(id=point_id))
        await db_session.commit()
    await db_session.rollback()
    with pytest.raises(IntegrityError, match='SRID'):
        await db_session.execute(text('UPDATE levantamientos SET srid=6362 WHERE id=:id')
            .bindparams(bindparam('id', type_=DBUUID(as_uuid=True))), dict(id=survey_id))
        await db_session.commit()
    await db_session.rollback()


@pytest.mark.asyncio
async def test_cut_and_fill_budget_keeps_source_for_both_lines_and_exports(db_session, tenant_a_user, monkeypatch):
    from io import BytesIO
    from openpyxl import load_workbook
    tenant, user = tenant_a_user
    existing, _ = await surfaces(db_session, tenant, user)
    # Saved non-planar triangles have both cut and fill against this plane.
    svc = TopografiaService(db_session, tenant.id)
    reconstruct = PresupuestoService._reconstruir_partidas_desde_db
    async def inverted_rows(self, *args, **kwargs):
        result = await reconstruct(self, *args, **kwargs)
        result.partidas.sort(key=lambda p: -p.numero)
        return result
    monkeypatch.setattr(PresupuestoService, '_reconstruir_partidas_desde_db', inverted_rows)
    calc = await svc.calcular_volumen(superficie_existente_id=existing.id, elevacion_referencia=.25,
                                     creado_por_id=user.id)
    budget = await svc.generar_partida_movimiento_tierras(calculo_volumen_id=calc.id,
        expediente_id=existing.expediente_id, parametros_costeo=PARAMS, creado_por_id=user.id)
    assert len(budget.partidas) == 2
    assert {p.metadatos['topografia']['tipo'] for p in budget.partidas} == {'CORTE', 'TERRAPLEN'}
    assert all(p.metadatos['topografia']['calculo_id'] == str(calc.id) for p in budget.partidas)
    assert budget.creado_por_id == user.id
    evidence = budget.metadatos['topografia_evidencia']
    assert evidence == calc.evidencia
    from app.models.catalogo_apu import CatalogoAPU
    catalog = CatalogoAPU(tenant_id=tenant.id, clave=uuid4().hex, descripcion='Measured earthwork fixture',
        tipo='CONCEPTO', unidad='m3', precio_unitario=100, fuente='TEST_SYNTHETIC')
    db_session.add(catalog)
    await db_session.commit()
    presupuestos = PresupuestoService(db_session, tenant.id)
    for p in budget.partidas:
        role = 'CORTE' if p.numero == 1 else 'TERRAPLEN'
        assert p.metadatos['topografia']['tipo'] == role
        budget = await presupuestos.asignar_catalogo_partida(budget.id, existing.expediente_id, p.id, catalog.id, user.id)
    assert budget.estado == 'CALCULADO'
    await presupuestos.cambiar_estado(budget.id, existing.expediente_id, 'VALIDADO')
    await presupuestos.cambiar_estado(budget.id, existing.expediente_id, 'APROBADO')
    exported = load_workbook(BytesIO(await presupuestos.generar_excel(budget.id, existing.expediente_id)))
    assert exported.active.title == 'Topografía'
    assert exported.active['B2'].value == str(calc.id)
    assert exported.active['B3'].value == evidence['sha256']
    assert budget.metadatos['topografia_evidencia'] == evidence
    for operation in [lambda: presupuestos.actualizar_cantidad_partida(budget.id, existing.expediente_id, budget.partidas[0].id, 3),
                      lambda: presupuestos.eliminar_partida(budget.id, existing.expediente_id, budget.partidas[0].id)]:
        with pytest.raises(MegalodonException, match='topográfic|movimiento de tierras'):
            await operation()


@pytest.mark.asyncio
async def test_partial_overlap_keeps_measured_cost_but_blocks_approval(db_session, tenant_a_user):
    from io import BytesIO
    from openpyxl import load_workbook
    from pypdf import PdfReader
    tenant, user = tenant_a_user
    existing, project = await surfaces(db_session, tenant, user)
    project.malla_vertices = [0,0,0, 1,0,0, 0,1,0]
    project.malla_caras = [0,1,2]
    # Deliberately misleading stored statistics must not certify actual coverage.
    project.area_plan_m2 = 1
    await db_session.commit()
    svc = TopografiaService(db_session, tenant.id)
    calc = await svc.calcular_volumen(superficie_existente_id=existing.id, superficie_proyecto_id=project.id)
    assert calc.evidencia['cobertura']['area_existente_pendiente_m2'] == .5
    assert calc.evidencia['cobertura']['completa'] is False
    budget = await svc.generar_partida_movimiento_tierras(calculo_volumen_id=calc.id,
        expediente_id=existing.expediente_id, parametros_costeo=PARAMS)
    from app.models.catalogo_apu import CatalogoAPU
    catalog = CatalogoAPU(tenant_id=tenant.id, clave=uuid4().hex, descripcion='Partial earthwork fixture',
        tipo='CONCEPTO', unidad='m3', precio_unitario=100, fuente='TEST_SYNTHETIC')
    db_session.add(catalog)
    await db_session.commit()
    presupuestos = PresupuestoService(db_session, tenant.id)
    for p in budget.partidas:
        await presupuestos.asignar_catalogo_partida(budget.id, existing.expediente_id, p.id, catalog.id, user.id)
    budget = await presupuestos.recalcular(budget.id, existing.expediente_id)
    assert budget.monto_total > 0
    assert budget.estado == 'BORRADOR'
    with pytest.raises(MegalodonException) as failure:
        await presupuestos.cambiar_estado(budget.id, existing.expediente_id, 'CALCULADO')
    assert failure.value.details['topografia_cobertura']['completa'] is False
    workbook = load_workbook(BytesIO(await presupuestos.generar_excel(budget.id, existing.expediente_id)))
    assert 'PARCIAL' in workbook.active['A1'].value
    assert workbook.active['B5'].value == .5
    pdf = PdfReader(BytesIO(await presupuestos.generar_pdf(budget.id, existing.expediente_id)))
    assert 'PARCIAL' in ''.join(page.extract_text() for page in pdf.pages)


@pytest.mark.asyncio
@pytest.mark.parametrize('change', ['mesh', 'quantity', 'reference', 'legacy', 'expediente'])
async def test_budget_rejects_changed_or_unproven_volume_sources(db_session, tenant_a_user, change):
    from tests.integration.test_closure_two_tenants import make_expediente
    tenant, user = tenant_a_user
    existing, _ = await surfaces(db_session, tenant, user)
    svc = TopografiaService(db_session, tenant.id)
    calc = await svc.calcular_volumen(superficie_existente_id=existing.id, elevacion_referencia=0)
    target = existing.expediente_id
    if change == 'mesh':
        existing.malla_caras = [0,1,2,0,2,3]
    elif change == 'quantity':
        calc.volumen_corte_m3 = 99
    elif change == 'reference':
        calc.elevacion_referencia = 50
    elif change == 'legacy':
        calc.evidencia = None
    else:
        other = make_expediente(tenant, user)
        db_session.add(other)
        await db_session.flush()
        target = other.id
    await db_session.commit()
    before = await db_session.scalar(select(func.count()).select_from(Presupuesto))
    with pytest.raises(MegalodonException):
        await svc.generar_partida_movimiento_tierras(calculo_volumen_id=calc.id,
            expediente_id=target, parametros_costeo=PARAMS)
    assert await db_session.scalar(select(func.count()).select_from(Presupuesto)) == before


@pytest.mark.asyncio
async def test_budget_source_write_failure_rolls_back_budget_lines_and_usage(db_session, tenant_a_user):
    from sqlalchemy.exc import DBAPIError
    from app.models.entitlements import TenantUso
    tenant, user = tenant_a_user
    tenant_id = tenant.id
    existing, _ = await surfaces(db_session, tenant, user)
    exp_id = existing.expediente_id
    svc = TopografiaService(db_session, tenant_id)
    calc = await svc.calcular_volumen(superficie_existente_id=existing.id, elevacion_referencia=.25)
    calc_id = calc.id
    before = await db_session.scalar(select(func.count()).select_from(Presupuesto).where(Presupuesto.expediente_id == exp_id))
    usage = await db_session.scalar(select(func.coalesce(func.sum(TenantUso.corridas_costeo),0)).where(TenantUso.tenant_id == tenant_id))
    sqlite = db_session.bind.dialect.name == 'sqlite'
    if sqlite:
        ddl = """CREATE TRIGGER reject_topo_source BEFORE UPDATE OF metadatos ON partidas
            WHEN json_extract(NEW.metadatos,'$.topografia') IS NOT NULL
            BEGIN SELECT RAISE(ABORT,'topo_source_rejected'); END"""
    else:
        await db_session.execute(text("""CREATE FUNCTION reject_topo_source() RETURNS trigger LANGUAGE plpgsql AS $$
            BEGIN IF NEW.metadatos ? 'topografia' THEN RAISE EXCEPTION 'topo_source_rejected'; END IF; RETURN NEW; END $$"""))
        ddl = "CREATE TRIGGER reject_topo_source BEFORE UPDATE OF metadatos ON partidas FOR EACH ROW EXECUTE FUNCTION reject_topo_source()"
    await db_session.execute(text(ddl))
    await db_session.commit()
    try:
        with pytest.raises(DBAPIError, match='topo_source_rejected'):
            await svc.generar_partida_movimiento_tierras(calculo_volumen_id=calc_id,
                expediente_id=exp_id, parametros_costeo=PARAMS)
        async with AsyncSessionLocalTest() as fresh:
            assert await fresh.scalar(select(func.count()).select_from(Presupuesto).where(Presupuesto.expediente_id == exp_id)) == before
            assert await fresh.scalar(select(func.coalesce(func.sum(TenantUso.corridas_costeo),0)).where(TenantUso.tenant_id == tenant_id)) == usage
            assert await fresh.get(CalculoVolumen, calc_id) is not None
    finally:
        await db_session.rollback()
        await db_session.execute(text('DROP TRIGGER reject_topo_source' + ('' if sqlite else ' ON partidas')))
        if not sqlite:
            await db_session.execute(text('DROP FUNCTION reject_topo_source()'))
        await db_session.commit()
