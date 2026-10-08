from app.engines.procurement.compiler import ProcurementArtifactCompiler
from io import BytesIO
from types import SimpleNamespace
import zipfile
from openpyxl import load_workbook
import pytest


def model():
    return {
        "identifier": "LIC-001",
        "title": "Obra de prueba real del motor",
        "jurisdiction_code": "FEDERAL",
        "revision": 1,
        "economic": {
            "budget_total": 110,
            "partidas": [{
                "numero": 1, "descripcion": "Movimiento de tierras", "unidad": "m3", "cantidad": 10,
                "precio_unitario": 11, "importe": 110,
                "conceptos": [{
                    "clave": "MT-001", "descripcion": "Excavación", "unidad": "m3", "cantidad": 10,
                    "costo_directo_unitario": 11,
                    "insumos": [{"clave": "MAT-01", "descripcion": "Material", "unidad": "m3", "cantidad": 10, "importe": 110}],
                }],
            }]
        },
        "schedule": {"duration_days": 10, "activities": [{"id": "A1", "wbs_code": "1.1", "name": "Excavación", "duration_days": 10, "budget": 110, "predecessors": []}]},
    }


def test_artifacts_are_reproducible():
    c = ProcurementArtifactCompiler()
    x1 = c.compile_economic_xlsx(model())
    x2 = c.compile_economic_xlsx(model())
    assert c.sha256(x1) == c.sha256(x2)
    p1 = c.compile_summary_pdf(model(), [])
    p2 = c.compile_summary_pdf(model(), [])
    assert c.sha256(p1) == c.sha256(p2)
    z1 = c.compile_submission_zip([{"path": "AE-02/v1.xlsx", "content": x1}], {"revision": 1})
    z2 = c.compile_submission_zip([{"path": "AE-02/v1.xlsx", "content": x2}], {"revision": 1})
    assert c.sha256(z1) == c.sha256(z2)


@pytest.mark.parametrize('method', ['compile_economic_xlsx', 'compile_apu_xlsx', 'compile_schedule_xlsx'])
def test_xlsx_is_reproducible_across_zip_clock_boundaries(monkeypatch, method):
    compiler = ProcurementArtifactCompiler()
    clock = SimpleNamespace(time=lambda: 0, localtime=lambda _: (2026, 1, 1, 0, 0, 0, 0, 1, 0))
    monkeypatch.setattr(zipfile, 'time', clock)
    before = getattr(compiler, method)(model())
    clock.localtime = lambda _: (2026, 10, 8, 23, 59, 59, 0, 1, 0)
    after = getattr(compiler, method)(model())
    assert before == after
    workbook = load_workbook(BytesIO(after))
    assert workbook.properties.modified.isoformat().startswith('2000-01-01')
    assert workbook.active.max_row > 1
