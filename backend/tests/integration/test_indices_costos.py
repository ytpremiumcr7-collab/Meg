"""Controlled fixtures are deliberately not official INEGI/CMIC observations."""
from io import BytesIO
from uuid import UUID, uuid4

import pytest
import pytest_asyncio
from openpyxl import load_workbook
from sqlalchemy import select, text
from sqlalchemy.exc import DBAPIError

from app.config import settings
from app.models.catalogo_conceptos import CatalogoFuente, InsumoCatalogo
from app.models.entitlements import PlanLimite
from app.models.user import User, UserRole
from app.services.auth_service import AuthService
from tests.conftest import _login
from tests.integration.test_closure_two_tenants import make_expediente

BASE = '/api/v1/indices-costos'


@pytest_asyncio.fixture
async def indices(async_client, db_session, tenant_a_user):
    tenant_a, user_a = tenant_a_user
    user_a.role = UserRole.SUPERADMIN.value
    admin = User(id=uuid4(), tenant_id=tenant_a.id, email=f'cost-{uuid4().hex}@test.invalid',
                 hashed_password=AuthService(db_session).hash_password('testpass123'),
                 full_name='Revisor de prueba', role=UserRole.ADMIN.value, is_active=True, is_verified=True)
    fuente = CatalogoFuente(id=uuid4(), nombre='CONTROLADO NO OFICIAL', tipo='CUSTOM',
                            vigencia_inicio='2020-01-01', vigencia_fin='2020-12-31', activo=True, moneda='MXN')
    expediente = make_expediente(tenant_a, admin)
    db_session.add_all([admin, fuente, expediente])
    if not await db_session.scalar(select(PlanLimite).where(PlanLimite.plan == 'FREE')):
        db_session.add(PlanLimite(plan='FREE', max_corridas_costeo_mes=3))
    await db_session.commit()
    global_auth, auth = await _login(async_client, user_a.email), await _login(async_client, admin.email)
    referencias = []
    for i, precio in enumerate(('20', '50')):
        insumo = InsumoCatalogo(id=uuid4(), fuente_id=fuente.id, clave=f'=1+{i}',
                               descripcion=f'Material controlado {i}', tipo='MATERIAL', unidad='kg',
                               precio_unitario=precio, incluye_iva=False, activo=True)
        db_session.add(insumo)
        await db_session.commit()
        evidencia = {'url': 'https://example.invalid/documento.pdf', 'sha256': str(i + 1) * 64,
                     'localizador': 'Tabla de prueba controlada, página 1'}
        series = await async_client.post(BASE + '/series', headers=global_auth, json={
            'codigo': f'TEST-{uuid4().hex}', 'version_metodologia': 'CONTROLADA-1',
            'nombre': 'Material sintético para pruebas', 'alcance': 'MATERIAL', 'region': 'NACIONAL',
            'moneda': 'MXN', 'incluye_iva': False, 'condiciones_precio': 'Precio controlado sin IVA y sin flete',
            'periodo_referencia': 'enero 2020=100', 'evidencia': evidencia,
        })
        assert series.status_code == 201, series.text
        serie = series.json()
        obs_ids = []
        for mes, valor, publicado, doc in [('2020-01-01', '100', '2020-02-10', 'a'),
                                           ('2020-08-01', '110' if i == 0 else '90', '2020-09-10', 'b')]:
            response = await async_client.post(BASE + '/observaciones', headers=global_auth, json={
                'serie_id': serie['id'], 'medida': 'NIVEL', 'mes': mes, 'valor': valor,
                'publicado_el': publicado, 'evidencia': {**evidencia, 'sha256': doc * 64},
            })
            assert response.status_code == 201, response.text
            obs_ids.append(response.json()['id'])
        response = await async_client.post(BASE + '/vinculos', headers=auth, json={
            'insumo_id': str(insumo.id), 'serie_id': serie['id'], 'mes_base': '2020-01-01',
            'region': 'NACIONAL', 'fundamento': 'Correspondencia revisada sólo para el ensayo controlado',
            'evidencia': evidencia,
        })
        assert response.status_code == 201, response.text
        referencias.append({'vinculo_id': response.json()['id'], 'observacion_base_id': obs_ids[0],
                            'observacion_destino_id': obs_ids[1]})
    return {'auth': auth, 'global_auth': global_auth, 'refs': referencias, 'expediente_id': str(expediente.id)}


@pytest.mark.asyncio
async def test_indices_apu_guardado_recalculo_excel_y_retiro(async_client, indices, db_session):
    refs, auth = indices['refs'], indices['auth']
    for referencia, precio in zip(refs, ('22.00', '45.00'), strict=True):
        response = await async_client.post(BASE + '/calcular', headers=auth, json=referencia)
        assert response.status_code == 200, response.text
        assert response.json()['precio_actualizado'] == precio
    root = f"/api/v1/presupuestos/{indices['expediente_id']}/presupuestos"
    body = {'nombre': 'APU indexado controlado', 'parametros_costeo': {
        'factor_indirecto': 0, 'factor_utilidad': 0, 'factor_impuesto': 0, 'factor_riesgo': 0,
        'fuente': 'CAPTURA_USUARIO', 'referencia': 'Datos sintéticos de prueba',
    }, 'partidas': [{'numero': 1, 'descripcion': '=1+1', 'unidad': '=SUM(A1)', 'cantidad': 2,
                     'insumos': [{'clave': 'dato cliente', 'descripcion': 'No autorizado', 'unidad': 'x',
                                  'tipo': 'MANO_OBRA', 'cantidad': cantidad, 'precio_unitario': 999,
                                  'actualizacion_precio': ref} for cantidad, ref in zip((2, 3), refs, strict=True)]}]}
    response = await async_client.post(root, headers=auth, json=body)
    assert response.status_code == 200, response.text
    presupuesto = response.json()
    assert presupuesto['monto_directo'] == 358
    insumos = presupuesto['partidas'][0]['conceptos'][0]['insumos']
    assert [i['precio_unitario'] for i in insumos] == [22, 45]
    assert all(i['tipo'] == 'MATERIAL' and i['unidad'] == 'kg' for i in insumos)
    assert len(presupuesto['metadatos']['actualizaciones_indices']) == 2
    old_sha = insumos[0]['actualizacion_precio']['sha256']
    # A later corrected publication and withdrawal must not reprice an issued budget.
    original = insumos[0]['actualizacion_precio']['destino']
    response = await async_client.post(BASE + '/observaciones', headers=indices['global_auth'], json={
        'serie_id': original['serie_id'], 'medida': 'NIVEL', 'mes': original['mes'], 'valor': '115',
        'publicado_el': '2020-10-01', 'evidencia': {**original['evidencia'], 'sha256': 'c' * 64},
    })
    assert response.status_code == 201, response.text
    nuevo = await async_client.post(BASE + '/calcular', headers=auth, json={
        **refs[0], 'observacion_destino_id': response.json()['id']})
    assert nuevo.json()['precio_actualizado'] == '23.00'
    retiro = await async_client.post(BASE + '/retiros', headers=indices['global_auth'], json={
        'observacion_id': refs[0]['observacion_destino_id'], 'motivo': 'Edición controlada reemplazada por revisión verificada'})
    assert retiro.status_code == 201, retiro.text
    rechazado = await async_client.post(BASE + '/calcular', headers=auth, json=refs[0])
    assert rechazado.status_code == 409, rechazado.text
    assert (await async_client.post(root, headers=auth, json=body)).status_code == 409
    recalc = await async_client.post(root + f"/{presupuesto['id']}/recalcular", headers=auth)
    assert recalc.status_code == 200, recalc.text
    assert recalc.json()['monto_total'] == 358
    assert recalc.json()['partidas'][0]['conceptos'][0]['insumos'][0]['actualizacion_precio']['sha256'] == old_sha
    excel = await async_client.get(root + f"/{presupuesto['id']}/excel", headers=auth)
    assert excel.status_code == 200, excel.text
    wb = load_workbook(BytesIO(excel.content), data_only=True)
    assert any(row[-1] == 358 for row in wb['Presupuesto'].iter_rows(values_only=True))
    assert wb['Actualizacion materiales']['F2'].value == '22.00'
    assert wb['Actualizacion materiales']['K2'].value == old_sha
    assert wb['Presupuesto']['B2'].data_type == 's'
    assert wb['Presupuesto']['C2'].data_type == 's'
    assert wb['Presupuesto']['F2'].data_type == 'n'
    assert wb['Actualizacion materiales']['A2'].data_type == 's'
    # Correcting a transcription preserves the real document hash.
    correction = await async_client.post(BASE + '/observaciones', headers=indices['global_auth'], json={
        'serie_id': original['serie_id'], 'medida': 'NIVEL', 'mes': original['mes'], 'valor': '101',
        'publicado_el': original['publicado_el'], 'evidencia': original['evidencia'],
        'revision_captura': 2, 'sustituye_id': original['id'],
    })
    assert correction.status_code == 201, correction.text
    assert correction.json()['documento_sha256'] == original['documento_sha256']
    corrected = await async_client.post(BASE + '/calcular', headers=auth, json={
        **refs[0], 'observacion_destino_id': correction.json()['id']})
    assert corrected.json()['precio_actualizado'] == '20.20'
    assert corrected.json()['destino']['sustituye_id'] == original['id']
    from app.models.presupuesto import Insumo
    guardado = await db_session.get(Insumo, UUID(insumos[0]['id']))
    for tipo, unidad in (('MATERIAL', 'ton'), ('MANO_OBRA', 'kg')):
        guardado.tipo, guardado.unidad = tipo, unidad
        await db_session.commit()
        response = await async_client.get(root + f"/{presupuesto['id']}/excel", headers=auth)
        assert response.status_code == 400 and 'identidad o unidad' in response.text


@pytest.mark.asyncio
async def test_indices_aislamiento_roles_y_datos_incompatibles(async_client, indices, tenant_b_user):
    _, user_b = tenant_b_user
    otro = await _login(async_client, user_b.email)
    assert (await async_client.get(BASE + '/vinculos', headers=otro)).json() == []
    assert (await async_client.post(BASE + '/calcular', headers=otro, json=indices['refs'][0])).status_code == 404
    retiro = await async_client.post(BASE + '/retiros', headers=otro, json={
        'vinculo_id': indices['refs'][0]['vinculo_id'], 'motivo': 'Intento controlado desde otro tenant'})
    assert retiro.status_code == 403
    response = await async_client.post(BASE + '/calcular', headers=indices['auth'], json={
        **indices['refs'][0], 'observacion_destino_id': indices['refs'][1]['observacion_destino_id']})
    assert response.status_code == 400
    response = await async_client.post(BASE + '/calcular', headers=indices['auth'], json={
        **indices['refs'][0], 'observacion_base_id': indices['refs'][0]['observacion_destino_id']})
    assert response.status_code == 400
    s = (await async_client.post(BASE + '/calcular', headers=indices['auth'], json=indices['refs'][0])).json()
    obs = {'serie_id': s['serie']['id'], 'medida': 'NIVEL', 'mes': '2020-01-01', 'valor': '100',
           'publicado_el': '2020-02-10', 'evidencia': s['base']['evidencia']}
    assert (await async_client.post(BASE + '/observaciones', headers=indices['auth'], json=obs)).status_code == 403
    for cambios in ({'medida': 'ANUAL'}, {'valor': 'NaN'},
                    {'evidencia': {**obs['evidencia'], 'url': 'https://example.invalid/doc?token=secret'}}):
        response = await async_client.post(BASE + '/observaciones', headers=indices['global_auth'], json={**obs, **cambios})
        assert response.status_code == 422, response.text


@pytest.mark.asyncio
@pytest.mark.skipif(not settings.database_async_url.startswith('postgresql'), reason='PostgreSQL trigger contract')
async def test_evidencia_inmutable_con_sql_directo(indices, db_session):
    for sentencia in ("UPDATE observaciones_indices_costos SET valor=150 WHERE id=:id",
                       "DELETE FROM observaciones_indices_costos WHERE id=:id"):
        with pytest.raises(DBAPIError):
            async with db_session.begin_nested():
                await db_session.execute(text(sentencia), {'id': indices['refs'][0]['observacion_base_id']})


@pytest.mark.asyncio
async def test_rechaza_moneda_desconocida_dolares_iva_y_salarios(async_client, indices, db_session):
    calculated = (await async_client.post(BASE + '/calcular', headers=indices['auth'], json=indices['refs'][0])).json()
    original = calculated['vinculo']['insumo_original']
    fuente = await db_session.get(CatalogoFuente, UUID(original['fuente_id']))
    insumo = await db_session.get(InsumoCatalogo, UUID(original['id']))
    payload = {'insumo_id': original['id'], 'serie_id': calculated['serie']['id'],
               'mes_base': '2020-01-01', 'region': 'NACIONAL',
               'fundamento': 'Intento controlado con condiciones incompatibles', 'evidencia': calculated['base']['evidencia']}
    for moneda, tipo, iva in ((None, 'MATERIAL', False), ('USD', 'MATERIAL', False),
                              ('MXN', 'MANO_OBRA', False), ('MXN', 'MATERIAL', True)):
        fuente.moneda, insumo.tipo, insumo.incluye_iva = moneda, tipo, iva
        await db_session.commit()
        response = await async_client.post(BASE + '/vinculos', headers=indices['auth'], json=payload)
        assert response.status_code == 400, response.text
    fuente.moneda, insumo.tipo, insumo.incluye_iva = None, 'MATERIAL', False
    await db_session.commit()
    acreditacion = {'fuente_id': str(fuente.id), 'moneda': 'MXN', 'evidencia': payload['evidencia']}
    response = await async_client.post(BASE + '/acreditaciones-moneda', headers=indices['auth'], json=acreditacion)
    assert response.status_code == 403
    response = await async_client.post(BASE + '/acreditaciones-moneda', headers=indices['global_auth'], json=acreditacion)
    assert response.status_code == 201, response.text
    assert response.json()['moneda_evidencia']['sha256'] == payload['evidencia']['sha256']
    assert (await async_client.post(BASE + '/acreditaciones-moneda', headers=indices['global_auth'], json=acreditacion)).status_code == 409
    response = await async_client.post(BASE + '/vinculos', headers=indices['auth'], json=payload)
    assert response.status_code == 201, response.text


@pytest.mark.asyncio
@pytest.mark.skipif(not settings.database_async_url.startswith('postgresql'), reason='PostgreSQL concurrent aggregate writes')
async def test_eliminaciones_concurrentes_preservan_manifiesto(async_client, indices, db_session):
    import asyncio

    from sqlalchemy.ext.asyncio import async_sessionmaker

    from app.services.presupuesto_service import PresupuestoService
    root = f"/api/v1/presupuestos/{indices['expediente_id']}/presupuestos"
    body = {'nombre': 'Eliminaciones simultáneas', 'parametros_costeo': {
        'factor_indirecto': 0, 'factor_utilidad': 0, 'factor_impuesto': 0, 'factor_riesgo': 0,
        'fuente': 'CAPTURA_USUARIO', 'referencia': 'Ensayo de dos transacciones PostgreSQL',
    }, 'partidas': [{'numero': n, 'descripcion': 'Material', 'unidad': 'kg', 'cantidad': 1,
                    'insumos': [{'cantidad': 1, 'actualizacion_precio': ref}]}
                   for n, ref in enumerate(indices['refs'], 1)]}
    response = await async_client.post(root, headers=indices['auth'], json=body)
    assert response.status_code == 200, response.text
    presupuesto = response.json()
    tenant = UUID((await async_client.post(BASE + '/calcular', headers=indices['auth'], json=indices['refs'][0])).json()['vinculo']['tenant_id'])
    maker = async_sessionmaker(db_session.bind, expire_on_commit=False)
    async def eliminar(id):
        async with maker() as session:
            service = PresupuestoService(session, tenant)
            return await service.eliminar_partida(UUID(presupuesto['id']), UUID(indices['expediente_id']), UUID(id))
    await asyncio.wait_for(asyncio.gather(*(eliminar(p['id']) for p in presupuesto['partidas'])), timeout=15)
    response = await async_client.get(root, headers=indices['auth'])
    assert response.status_code == 200, response.text
    guardado = next(p for p in response.json() if p['id'] == presupuesto['id'])
    assert guardado['partidas'] == []
    assert guardado['metadatos']['actualizaciones_indices'] == {}
    assert guardado['monto_total'] == 0


@pytest.mark.asyncio
@pytest.mark.skipif(not settings.database_async_url.startswith('postgresql'), reason='PostgreSQL shared export lock')
async def test_exportacion_coherente_durante_eliminacion(async_client, indices, db_session):
    import asyncio

    from sqlalchemy.ext.asyncio import async_sessionmaker

    from app.services.presupuesto_service import PresupuestoService
    root = f"/api/v1/presupuestos/{indices['expediente_id']}/presupuestos"
    body = {'nombre': 'Exportación simultánea', 'parametros_costeo': {
        'factor_indirecto': 0, 'factor_utilidad': 0, 'factor_impuesto': 0, 'factor_riesgo': 0,
        'fuente': 'CAPTURA_USUARIO', 'referencia': 'Ensayo de exportación PostgreSQL',
    }, 'partidas': [{'numero': 1, 'descripcion': 'Material', 'unidad': 'kg', 'cantidad': 1,
                    'insumos': [{'cantidad': 1, 'actualizacion_precio': indices['refs'][0]}]}]}
    response = await async_client.post(root, headers=indices['auth'], json=body)
    assert response.status_code == 200, response.text
    presupuesto = response.json()
    tenant = UUID((await async_client.post(BASE + '/calcular', headers=indices['auth'], json=indices['refs'][0])).json()['vinculo']['tenant_id'])
    maker = async_sessionmaker(db_session.bind, expire_on_commit=False)
    presupuesto_id, expediente_id = UUID(presupuesto['id']), UUID(indices['expediente_id'])
    async with maker() as lectura, maker() as escritura:
        lector = PresupuestoService(lectura, tenant)
        await lector._reconstruir_partidas_desde_db(presupuesto_id, expediente_id)
        pid = await escritura.scalar(text('SELECT pg_backend_pid()'))
        tarea = asyncio.create_task(PresupuestoService(escritura, tenant).eliminar_partida(
            presupuesto_id, expediente_id, UUID(presupuesto['partidas'][0]['id'])))
        try:
            async def esperar_bloqueo():
                while not await lectura.scalar(text('SELECT EXISTS(SELECT 1 FROM pg_locks WHERE pid=:pid AND NOT granted)'), {'pid': pid}):
                    assert not tarea.done(), 'El escritor modificó el agregado mientras se estaba exportando'
                    await asyncio.sleep(0.01)
            await asyncio.wait_for(esperar_bloqueo(), timeout=5)
            excel = await lector.generar_excel(presupuesto_id, expediente_id)
            wb = load_workbook(BytesIO(excel), data_only=True)
            assert wb['Presupuesto']['F2'].value == 22
            assert wb['Actualizacion materiales']['F2'].value == '22.00'
            await lectura.rollback()
            resultado = await asyncio.wait_for(tarea, timeout=5)
            assert resultado.monto_total == 0
        finally:
            await lectura.rollback()
            if not tarea.done():
                tarea.cancel()
                await asyncio.gather(tarea, return_exceptions=True)
