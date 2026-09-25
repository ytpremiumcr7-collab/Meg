from __future__ import annotations

import io
import json
import zipfile
from datetime import datetime
from hashlib import sha256
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas


class ProcurementArtifactCompiler:
    """Compiles canonical procurement data into real, reproducible artifacts.

    The compiler algorithms are intentionally code-owned. Which external format
    invokes which algorithm is DB-owned and resolved through this allow-list;
    arbitrary method names from the database are never executed.
    """

    _COMPILER_REGISTRY = {
        "compile_proposal_letter_xlsx": "compile_proposal_letter_xlsx",
        "compile_economic_xlsx": "compile_economic_xlsx",
        "compile_apu_xlsx": "compile_apu_xlsx",
        "compile_lump_sum_xlsx": "compile_lump_sum_xlsx",
        "compile_mixed_contract_xlsx": "compile_mixed_contract_xlsx",
        "compile_schedule_xlsx": "compile_schedule_xlsx",
        "compile_annex_pdf": "compile_annex_pdf",
        "compile_fasar_xlsx": "compile_fasar_xlsx",
        "compile_indirects_xlsx": "compile_indirects_xlsx",
        "compile_financing_xlsx": "compile_financing_xlsx",
        "compile_profit_xlsx": "compile_profit_xlsx",
        "compile_material_schedule_xlsx": "compile_material_schedule_xlsx",
        "compile_labor_schedule_xlsx": "compile_labor_schedule_xlsx",
        "compile_equipment_schedule_xlsx": "compile_equipment_schedule_xlsx",
        "compile_explosion_xlsx": "compile_explosion_xlsx",
        "compile_budget_amount_xlsx": "compile_budget_amount_xlsx",
    }

    def compile_from_binding(self, binding: dict[str, Any], model: dict[str, Any]) -> bytes:
        compiler_name = str(binding.get("compiler") or "").strip()
        method_name = self._COMPILER_REGISTRY.get(compiler_name)
        if method_name is None:
            raise ValueError(f"Compiler no permitido o inexistente para el formato {binding.get('format_code')}: {compiler_name!r}")
        method = getattr(self, method_name, None)
        if method is None or not callable(method):
            raise ValueError(f"Implementación del compiler ausente: {compiler_name}")
        if method_name == "compile_annex_pdf":
            return method(str(binding.get("format_code")), str(binding.get("title") or binding.get("format_code")), model, binding.get("requirement"))
        return method(model)

    @staticmethod
    def _money(value: Any) -> float:
        return float(value or 0)

    def compile_proposal_letter_xlsx(self, model: dict[str, Any]) -> bytes:
        facts=model.get("facts", {})
        economic=model.get("economic", {})
        if not facts.get("authority") or not facts.get("object") or not economic.get("budget_total"):
            raise ValueError("La carta de propuesta requiere convocante, objeto y monto total determinados.")
        rows=[["Convocante",facts.get("authority")],["Procedimiento",facts.get("procedure_type")],["Tipo de contrato",facts.get("contract_type")],["Objeto",facts.get("object")],["Monto sin IVA",economic.get("budget_total")],["Vigencia",facts.get("proposal_validity_days")]]
        return self._xlsx_from_rows("Carta Propuesta", ["Campo","Valor"], rows, {"A":30,"B":90})

    def compile_economic_xlsx(self, model: dict[str, Any]) -> bytes:
        economic = model.get("economic", {})
        partidas = economic.get("partidas", [])
        if not partidas:
            raise ValueError("No existen partidas económicas para generar AE-02.")
        wb = Workbook()
        wb.properties.creator = "Megalodon"
        wb.properties.lastModifiedBy = "Megalodon"
        wb.properties.created = datetime(2000, 1, 1)
        wb.properties.modified = datetime(2000, 1, 1)
        ws = wb.active
        ws.title = "Catalogo"
        headers = ["No.", "Descripción", "Unidad", "Cantidad", "P.U.", "Importe"]
        ws.append(headers)
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center")
        for row in partidas:
            cantidad = self._money(row.get("cantidad"))
            pu = self._money(row.get("precio_unitario", row.get("precio_unitario_manual")))
            importe = self._money(row.get("importe", cantidad * pu))
            ws.append([row.get("numero"), row.get("descripcion"), row.get("unidad"), cantidad, pu, importe])
        ws.append(["", "", "", "", "TOTAL", self._money(economic.get("budget_total"))])
        for col, width in {"A": 8, "B": 60, "C": 14, "D": 16, "E": 16, "F": 18}.items():
            ws.column_dimensions[col].width = width
        out = io.BytesIO()
        wb.save(out)
        return out.getvalue()

    def compile_apu_xlsx(self, model: dict[str, Any]) -> bytes:
        partidas = model.get("economic", {}).get("partidas", [])
        concepts = [c for p in partidas for c in p.get("conceptos", [])]
        if not concepts:
            raise ValueError("No existen análisis de precio unitario para generar APU.")
        wb = Workbook()
        wb.properties.creator = "Megalodon"
        wb.properties.lastModifiedBy = "Megalodon"
        wb.properties.created = datetime(2000, 1, 1)
        wb.properties.modified = datetime(2000, 1, 1)
        ws = wb.active
        ws.title = "APU"
        ws.append(["Partida", "Concepto", "Descripción", "Unidad", "Cantidad", "Costo directo unitario"])
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for partida in partidas:
            for concept in partida.get("conceptos", []):
                ws.append([
                    partida.get("numero"), concept.get("clave"), concept.get("descripcion"),
                    concept.get("unidad"), self._money(concept.get("cantidad")),
                    self._money(concept.get("costo_directo_unitario")),
                ])
                for item in concept.get("insumos", []):
                    ws.append(["", f"  {item.get('clave')}", item.get("descripcion"), item.get("unidad"), self._money(item.get("cantidad")), self._money(item.get("importe"))])
        ws.column_dimensions["A"].width = 12
        ws.column_dimensions["B"].width = 24
        ws.column_dimensions["C"].width = 55
        ws.column_dimensions["D"].width = 14
        ws.column_dimensions["E"].width = 16
        ws.column_dimensions["F"].width = 24
        out = io.BytesIO(); wb.save(out); return out.getvalue()


    def compile_lump_sum_xlsx(self, model: dict[str, Any]) -> bytes:
        economic = model.get("economic", {})
        activities = economic.get("lump_sum_activities") or model.get("schedule", {}).get("activities", [])
        total = economic.get("lump_sum_total") or economic.get("budget_total")
        if not activities or total is None:
            raise ValueError("Precio alzado requiere actividades, monto total y trazabilidad.")
        rows = []
        total_value = self._money(total)
        for item in activities:
            rows.append([
                item.get("id"), item.get("wbs_code"), item.get("name"),
                self._money(item.get("budget")), item.get("start_date"), item.get("end_date"),
            ])
        return self._xlsx_from_rows(
            "Precio Alzado",
            ["ID", "WBS", "Actividad", "Importe", "Inicio", "Fin"],
            rows + [["", "", "TOTAL", total_value, "", ""]],
            {"A": 18, "B": 18, "C": 60, "D": 20, "E": 16, "F": 16},
        )

    def compile_mixed_contract_xlsx(self, model: dict[str, Any]) -> bytes:
        components = model.get("economic", {}).get("mixed_components") or []
        if not components:
            raise ValueError("Contrato mixto requiere componentes a precios unitarios y precio alzado.")
        rows = []
        for component in components:
            basis = str(component.get("basis", "")).upper()
            if basis == "UNIT_PRICES":
                amount = component.get("budget_total") or sum(self._money(p.get("importe", 0)) for p in component.get("partidas", []))
            else:
                amount = component.get("total") or component.get("budget_total")
            if amount is None:
                raise ValueError("Cada componente mixto requiere importe determinable.")
            rows.append([component.get("code"), basis, component.get("description"), self._money(amount)])
        total = sum(self._money(r[3]) for r in rows)
        rows.append(["", "TOTAL", "Contrato mixto", total])
        return self._xlsx_from_rows("Contrato Mixto", ["Código", "Base", "Descripción", "Importe"], rows, {"A": 18, "B": 20, "C": 60, "D": 20})

    def compile_schedule_xlsx(self, model: dict[str, Any]) -> bytes:
        schedule = model.get("schedule", {})
        activities = schedule.get("activities", [])
        if not activities:
            raise ValueError("No existen actividades para generar el programa de obra.")
        wb = Workbook()
        wb.properties.creator = "Megalodon"
        wb.properties.lastModifiedBy = "Megalodon"
        wb.properties.created = datetime(2000, 1, 1)
        wb.properties.modified = datetime(2000, 1, 1)
        ws = wb.active
        ws.title = "Programa"
        ws.append(["ID", "WBS", "Actividad", "Duración (días)", "Presupuesto", "Predecesoras"])
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for a in activities:
            ws.append([a.get("id"), a.get("wbs_code"), a.get("name"), self._money(a.get("duration_days")), self._money(a.get("budget")), ",".join(map(str, a.get("predecessors", [])))])
        ws.append(["", "", "DURACIÓN TOTAL", self._money(schedule.get("duration_days")), "", ""])
        for col, width in {"A": 18, "B": 18, "C": 58, "D": 18, "E": 18, "F": 30}.items():
            ws.column_dimensions[col].width = width
        out = io.BytesIO(); wb.save(out); return out.getvalue()

    def compile_summary_pdf(self, model: dict[str, Any], findings: list[dict[str, Any]]) -> bytes:
        if not model.get("identifier"):
            raise ValueError("El modelo canónico requiere identifier para generar el reporte.")
        out = io.BytesIO()
        c = canvas.Canvas(out, pagesize=letter, invariant=1)
        c.setCreator("Megalodon")
        c.setTitle(str(model.get("title", "Tender Package"))[:200])
        y = 750
        c.setFont("Helvetica-Bold", 14)
        c.drawString(50, y, str(model.get("title", "Tender Package"))[:100]); y -= 28
        c.setFont("Helvetica", 9)
        for label, value in ((
            ("Identificador", model.get("identifier")),
            ("Presupuesto", model.get("economic", {}).get("budget_total")),
            ("Jurisdicción", model.get("jurisdiction_code")),
            ("Revisión", model.get("revision")),
            ("Modelo", ProcurementArtifactCompiler.sha256(json.dumps(model, sort_keys=True, separators=(",", ":"), default=str).encode("utf-8"))),
        )):
            c.drawString(50, y, f"{label}: {value}"); y -= 16
        y -= 8; c.setFont("Helvetica-Bold", 10); c.drawString(50, y, "Validaciones"); y -= 18; c.setFont("Helvetica", 9)
        for f in findings:
            line = f"[{f.get('severity')}] {f.get('code')}: {f.get('message')}"
            c.drawString(50, y, line[:125]); y -= 14
            if y < 60:
                c.showPage(); y = 750; c.setFont("Helvetica", 9)
        c.save(); return out.getvalue()


    def compile_annex_pdf(self, code: str, title: str, model: dict[str, Any], requirement: dict[str, Any] | None = None) -> bytes:
        """Render a deterministic working annex from canonical data.

        External format codes are metadata selected by DB. This algorithm never
        interprets those codes; DB canonical_paths select the payload when present.
        """
        if not code or not title:
            raise ValueError("El anexo requiere código y título.")
        selected: dict[str, Any] = {}
        for path in (requirement or {}).get("canonical_paths") or []:
            cur: Any = model
            for part in str(path).split("."):
                if not isinstance(cur, dict):
                    cur = None
                    break
                cur = cur.get(part)
            if cur not in (None, "", [], {}):
                selected[str(path)] = cur
        payload: Any = selected or {
            "facts": model.get("facts", {}),
            "technical": model.get("technical", {}),
            "economic": model.get("economic", {}),
            "schedule": model.get("schedule", {}),
            "bidder": model.get("bidder", {}),
        }
        def _has_data(value: Any) -> bool:
            if value is None or value == "": return False
            if isinstance(value, dict): return any(_has_data(v) for v in value.values())
            if isinstance(value, (list, tuple, set)): return any(_has_data(v) for v in value)
            if isinstance(value, (int, float)): return value != 0
            return True
        if not _has_data(payload):
            raise ValueError(f"No existe información canónica suficiente para {code}; se requiere evidencia/modelo antes de generar el artefacto.")
        out = io.BytesIO()
        c = canvas.Canvas(out, pagesize=letter, invariant=1)
        c.setCreator("Megalodon"); c.setTitle(str(title)[:200])
        y = 760
        c.setFont("Helvetica-Bold", 14); c.drawString(45, y, str(title)[:100]); y -= 20
        c.setFont("Helvetica", 8); c.drawString(45, y, f"Código: {code} | Expediente: {model.get('identifier')} | Revisión: {model.get('revision', 1)}"); y -= 18
        facts = model.get("facts", {})
        c.setFont("Helvetica", 9); c.drawString(45, y, f"Jurisdicción: {model.get('jurisdiction_code')} | Tipo de contrato: {facts.get('contract_type')}"); y -= 22
        for line in json.dumps(payload, ensure_ascii=False, indent=2, sort_keys=True, default=str).splitlines():
            if y < 55:
                c.showPage(); y = 760; c.setFont("Helvetica", 9)
            c.drawString(45, y, line[:120]); y -= 11
        c.save()
        return out.getvalue()


    def compile_indirects_xlsx(self, model: dict[str, Any]) -> bytes:
        d = model.get("economic", {}).get("indirect_costs") or {}
        items=d.get("items") or []
        if not items: raise ValueError("Los costos indirectos requieren desglose documentado; no se admite porcentaje mágico.")
        rows=[[i.get("concept"), i.get("basis"), i.get("amount")] for i in items]
        rows.append(["TOTAL", "", d.get("total")])
        return self._xlsx_from_rows("Indirectos", ["Concepto","Base","Importe"], rows, {"A":45,"B":28,"C":20})

    def compile_financing_xlsx(self, model: dict[str, Any]) -> bytes:
        f=model.get("economic", {}).get("financing") or {}
        flows=f.get("cash_flow") or []
        if not flows or f.get("factor") is None: raise ValueError("Financiamiento requiere flujo de efectivo y factor derivado.")
        rows=[[x.get("period"), x.get("inflow"), x.get("outflow"), x.get("net")] for x in flows]
        rows.append(["FACTOR", "", "", f.get("factor")])
        return self._xlsx_from_rows("Financiamiento", ["Periodo","Ingreso","Egreso","Neto"], rows, {"A":16,"B":18,"C":18,"D":18})

    def compile_profit_xlsx(self, model: dict[str, Any]) -> bytes:
        p=model.get("economic", {}).get("profit_detail") or {}
        if p.get("amount") is None or p.get("basis") is None: raise ValueError("Utilidad requiere monto y base de cálculo trazables.")
        return self._xlsx_from_rows("Utilidad", ["Variable","Valor"], [["Base",p["basis"]],["Porcentaje",p.get("rate")],["Importe",p["amount"]],["Criterio",p.get("criterion")]], {"A":28,"B":40})

    def compile_material_schedule_xlsx(self, model: dict[str, Any]) -> bytes:
        items=model.get("economic", {}).get("material_schedule") or []
        if not items: raise ValueError("Programa de suministro de materiales requiere cantidades y periodos.")
        return self._xlsx_from_rows("Suministro", ["Clave","Material","Unidad","Cantidad","Periodo"], [[x.get("key"),x.get("name"),x.get("unit"),x.get("quantity"),x.get("period")] for x in items], {"A":16,"B":48,"C":14,"D":16,"E":18})

    def compile_labor_schedule_xlsx(self, model: dict[str, Any]) -> bytes:
        items=model.get("economic", {}).get("labor_schedule") or []
        if not items: raise ValueError("Programa de mano de obra requiere personal/categorías y periodos.")
        return self._xlsx_from_rows("Programa MO", ["Categoría","Cantidad","Horas/Día","Periodo"], [[x.get("category"),x.get("quantity"),x.get("hours_per_day"),x.get("period")] for x in items], {"A":40,"B":16,"C":16,"D":18})

    def compile_equipment_schedule_xlsx(self, model: dict[str, Any]) -> bytes:
        items=model.get("economic", {}).get("equipment_schedule") or []
        if not items: raise ValueError("Programa de equipo requiere maquinaria, horas/cantidad y periodos.")
        return self._xlsx_from_rows("Programa Equipo", ["Equipo","Cantidad","Horas","Periodo"], [[x.get("name"),x.get("quantity"),x.get("hours"),x.get("period")] for x in items], {"A":40,"B":16,"C":16,"D":18})

    def compile_explosion_xlsx(self, model: dict[str, Any]) -> bytes:
        aggregated: dict[tuple[str,str], float] = {}
        for p in model.get("economic", {}).get("partidas", []):
            for c in p.get("conceptos", []):
                for item in c.get("insumos", []):
                    key=(str(item.get("clave")), str(item.get("unidad")))
                    aggregated[key]=aggregated.get(key,0.0)+self._money(item.get("cantidad")) * self._money(c.get("cantidad",1))
        if not aggregated: raise ValueError("La explosión de insumos requiere conceptos con insumos reales.")
        rows=[[k[0],k[1],v] for k,v in sorted(aggregated.items())]
        return self._xlsx_from_rows("Explosion Insumos", ["Clave","Unidad","Cantidad Total"], rows, {"A":22,"B":14,"C":20})

    def compile_budget_amount_xlsx(self, model: dict[str, Any]) -> bytes:
        total=model.get("economic", {}).get("budget_total")
        words=model.get("economic", {}).get("amount_in_words")
        if not total or not words: raise ValueError("AE-12 requiere monto total y su representación en letra provista/calculada por la capa económica.")
        return self._xlsx_from_rows("Presupuesto", ["Campo","Valor"], [["Total",total],["En letra",words]], {"A":24,"B":80})

    def compile_submission_zip(self, artifacts: list[dict[str, Any]], manifest: dict[str, Any]) -> bytes:
        if not artifacts:
            raise ValueError("No hay artefactos para construir el paquete de presentación.")
        out = io.BytesIO()
        fixed_date = (2000, 1, 1, 0, 0, 0)
        with zipfile.ZipFile(out, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9) as archive:
            for item in sorted(artifacts, key=lambda x: str(x["path"])):
                path = str(item["path"]).lstrip("/")
                content = item["content"]
                if not isinstance(content, (bytes, bytearray)):
                    raise TypeError(f"Contenido inválido para artefacto {path}")
                info = zipfile.ZipInfo(path, date_time=fixed_date)
                info.compress_type = zipfile.ZIP_DEFLATED
                info.create_system = 3
                info.external_attr = 0o100644 << 16
                archive.writestr(info, bytes(content))
            manifest_bytes = json.dumps(manifest, ensure_ascii=False, indent=2, sort_keys=True).encode("utf-8")
            info = zipfile.ZipInfo("MANIFEST.json", date_time=fixed_date)
            info.compress_type = zipfile.ZIP_DEFLATED
            info.create_system = 3
            info.external_attr = 0o100644 << 16
            archive.writestr(info, manifest_bytes)
        return out.getvalue()

    @staticmethod
    def sha256(content: bytes) -> str:
        return sha256(content).hexdigest()
