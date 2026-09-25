from __future__ import annotations

import io
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any


@dataclass(frozen=True)
class ExtractedSource:
    filename: str
    media_type: str
    text: str
    pages: int | None
    metadata: dict[str, Any]


@dataclass(frozen=True)
class ExtractedRequirementCandidate:
    code: str
    description: str
    category: str
    mandatory: bool | None
    source_reference: dict[str, Any]
    confidence: str = "STRUCTURED_SOURCE"


class TenderSourceExtractor:
    """Deterministic document-structure extraction for tender sources.

    No LLM, embeddings, semantic similarity, legal-corpus mining, or probabilistic
    requirement generation is used. The engine reads structures that physically exist
    in the supplied source and maps them through dependency/flow/format configuration.
    """

    MAX_TEXT_CHARS = 5_000_000
    _OFFICIAL_ID_PATTERNS = (
        re.compile(r"\b(?:REQUISITO|REQ(?:UISITO)?|R)\s*[-.:#]?\s*([A-Z0-9][A-Z0-9._/-]{0,30})\b", re.I),
        re.compile(r"\b(?:FORMATO|ANEXO|AP[ÉE]NDICE|APENDICE)\s*[-.:#]?\s*([A-Z0-9][A-Z0-9._/-]{0,30})\b", re.I),
        re.compile(r"^\s*(\d+(?:\.\d+){0,5})[.)]?\s+", re.I),
        re.compile(r"^\s*([A-Z]{1,4}[-_]?\d{1,6})[.)]?\s+", re.I),
    )
    _RELATION_PATTERNS = (
        ("MODIFICA", re.compile(r"\b(?:se\s+)?modifica(?:do|da|dos|das)?\b|\bmodificaci[oó]n\b", re.I)),
        ("SUSTITUYE", re.compile(r"\b(?:se\s+)?sustitu(?:ye|ir[aá]|ido|ida|idos|idas)\b|\ben\s+sustituci[oó]n\s+de\b", re.I)),
        ("ACLARA", re.compile(r"\b(?:se\s+)?aclara(?:ci[oó]n)?\b|\bprecisa(?:ci[oó]n)?\b", re.I)),
    )

    def extract(self, filename: str, media_type: str, content: bytes) -> ExtractedSource:
        if not content:
            raise ValueError("El documento fuente está vacío.")
        suffix = Path(filename).suffix.lower()
        if suffix == ".pdf" or media_type == "application/pdf":
            return self._pdf(filename, media_type, content)
        if suffix == ".docx" or media_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
            return self._docx(filename, media_type, content)
        if suffix in {".xlsx", ".xlsm"} or "spreadsheet" in media_type or "excel" in media_type:
            return self._xlsx(filename, media_type, content)
        if suffix in {".txt", ".csv", ".json", ".xml", ".md"} or media_type.startswith("text/") or "json" in media_type or "xml" in media_type:
            text = content.decode("utf-8-sig", errors="strict")
            return ExtractedSource(filename, media_type, text[: self.MAX_TEXT_CHARS], None, {"format": suffix.lstrip("."), "blocks": self._text_blocks(text)})
        raise ValueError(f"Formato documental no soportado sin un extractor explícito: {filename}")

    def detect_document_version(self, extracted: ExtractedSource, pattern: str | None = None) -> str | None:
        rx = pattern or r"(?:versi[oó]n|version|rev(?:isi[oó]n)?|revision)\s*[:#-]?\s*([A-Za-z0-9._-]+)"
        match = re.search(rx, extracted.text or "", re.I)
        return match.group(1).strip() if match else None

    def extract_requirement_candidates(
        self, extracted: ExtractedSource, *, source_id: str, source_hash: str,
        source_role: str, extraction_config: dict[str, Any] | None = None,
        format_code: str | None = None, physical_format: str | None = None,
        flow_code: str | None = None, dependency_code: str | None = None,
        document_version: str | None = None, source_revision: int | None = None,
    ) -> list[dict[str, Any]]:
        """Map only explicitly configured source structures into candidates.

        A candidate is a transcription/mapping result, not a legal inference. Mandatory
        status remains null unless the source itself contains an explicit configured field.
        """
        config = extraction_config or {}
        role_cfg = (config.get("source_roles") or {}).get(source_role)
        if not isinstance(role_cfg, dict):
            return []
        flows = role_cfg.get("flows") or {}
        flow_cfg = flows.get(str(flow_code or "").upper()) or flows.get("*")
        if isinstance(flow_cfg, dict):
            role_cfg = {**role_cfg, **flow_cfg}
        fmt = self._select_format(role_cfg, format_code, physical_format)
        sections = fmt.get("sections") or []
        if not sections:
            # Backward-compatible compact config, still deterministic.
            headings = fmt.get("headings") or []
            if headings:
                sections = [{"name": str(h), "headings": [str(h)], "item_pattern": fmt.get("item_pattern"), "category": fmt.get("category", "ADMINISTRATIVO")} for h in headings]
        blocks = list(extracted.metadata.get("blocks") or [])
        candidates: list[dict[str, Any]] = []
        for section_cfg in sections:
            if not isinstance(section_cfg, dict):
                continue
            candidates.extend(self._extract_section(
                blocks, section_cfg, source_id=source_id, source_hash=source_hash,
                source_role=source_role, flow_code=flow_code,
                dependency_code=dependency_code or config.get("dependency_code"),
                format_code=str(format_code or extracted.metadata.get("format") or "UNKNOWN").upper(),
                document_version=document_version, source_revision=source_revision, filename=extracted.filename,
            ))
        # Table contracts may be the complete extraction contract for spreadsheet/forms.
        candidates.extend(self._extract_tables(
            blocks, fmt, source_id=source_id, source_hash=source_hash,
            source_role=source_role, flow_code=flow_code,
            dependency_code=dependency_code or config.get("dependency_code"),
            format_code=str(format_code or extracted.metadata.get("format") or "UNKNOWN").upper(),
            document_version=document_version, source_revision=source_revision, filename=extracted.filename,
        ))
        return self._dedupe_candidates(candidates)

    def extract_source_relations(
        self, extracted: ExtractedSource, *, source_id: str, source_hash: str,
        source_role: str, dependency_code: str | None, flow_code: str | None,
        format_code: str | None = None,
        document_version: str | None = None,
        source_revision: int | None = None,
    ) -> list[dict[str, Any]]:
        """Extract explicit MODIFICA/SUSTITUYE/ACLARA statements.

        Targets are linked only when an explicit requirement/format identifier is present.
        Otherwise the relation is retained as unresolved and cannot alter an effective
        requirement automatically.
        """
        relations: list[dict[str, Any]] = []
        blocks = list(extracted.metadata.get("blocks") or [])
        grouped: list[dict[str, Any]] = []
        table_groups: dict[tuple[Any, Any, Any], list[dict[str, Any]]] = {}
        for block in blocks:
            if block.get("table") is not None and block.get("row") is not None:
                table_groups.setdefault((block.get("page"), block.get("table"), block.get("row")), []).append(block)
        grouped.extend({"text": " | ".join(str(x.get("text") or "").strip() for x in sorted(v, key=lambda z: str(z.get("cell"))) if str(x.get("text") or "").strip()), "anchor": v[0]} for v in table_groups.values())
        grouped.extend({"text": str(b.get("text") or "").strip(), "anchor": b} for b in blocks if b.get("table") is None or b.get("row") is None)
        for item in grouped:
            text = item["text"]
            block = item["anchor"]
            if not text:
                continue
            relation_type = next((kind for kind, rx in self._RELATION_PATTERNS if rx.search(text)), None)
            if not relation_type:
                continue
            target_id = self._explicit_identifier(text)
            relations.append({
                "relation_type": relation_type,
                "target_identifier": target_id,
                "resolved": bool(target_id),
                "source_reference": {
                    "source_id": source_id, "source_hash": source_hash,
                    "source_role": source_role, "document": extracted.filename,
                    "document_version": document_version, "revision": source_revision,
                    "page": block.get("page"), "section": block.get("section"),
                    "table": block.get("table"), "row": block.get("row"),
                    "cell": block.get("cell"), "field": block.get("field"),
                    "line_start": block.get("line"), "line_end": block.get("line"),
                    "text_original": text, "authority": "TENDER_SOURCE",
                    "extraction_method": "DETERMINISTIC_EXPLICIT_RELATION",
                },
                "dependency_code": dependency_code,
                "flow_code": flow_code,
                "format_code": format_code,
                "status": "MAPPED" if target_id else "REVIEW_REQUIRED",
            })
        return relations

    def _extract_section(self, blocks: list[dict[str, Any]], cfg: dict[str, Any], **ctx: Any) -> list[dict[str, Any]]:
        headings = [str(x).strip() for x in (cfg.get("headings") or [cfg.get("name")]) if str(x).strip()]
        stops = [str(x).strip() for x in (cfg.get("stop_headings") or []) if str(x).strip()]
        if not headings:
            return []
        heading_re = re.compile(r"^(?:#+\s*)?(?:" + "|".join(re.escape(x) for x in headings) + r")\s*:?(?:\s+.*)?$", re.I)
        stop_re = re.compile(r"^(?:#+\s*)?(?:" + "|".join(re.escape(x) for x in stops) + r")\s*:?(?:\s+.*)?$", re.I) if stops else None
        item_re = re.compile(str(cfg.get("item_pattern") or r"^\s*(?:[-•*]|\d+[.)]|[A-Za-z][.)])\s+(.+?)\s*$"))
        active = False
        section_name = None
        out: list[dict[str, Any]] = []
        ordinal = 0
        for idx, block in enumerate(blocks, start=1):
            raw = str(block.get("text") or "")
            line = raw.strip()
            if not line:
                continue
            if heading_re.match(line):
                active, section_name = True, line
                continue
            if active and stop_re and stop_re.match(line):
                active, section_name = False, None
                continue
            if not active:
                continue
            match = item_re.match(raw)
            if not match:
                continue
            description = (match.groupdict().get("description") if match.groupdict() else None) or (match.group(1) if match.lastindex else raw.strip())
            description = description.strip()
            if len(description) < 8:
                continue
            ordinal += 1
            official_id = self._explicit_identifier(raw) or self._explicit_identifier(description)
            code = official_id or f"SRC-{ctx['source_hash'][:10].upper()}-{ordinal:03d}"
            ref = self._location(block, ctx, section_name, raw)
            mandatory = self._explicit_mandatory(block, cfg)
            out.append({
                "code": code, "official_identifier": official_id,
                "description": description,
                "category": str(cfg.get("category") or "ADMINISTRATIVO").upper(),
                "mandatory": mandatory,
                "source_reference": ref,
                "mapping": {
                    "target": "TenderRequirement",
                    "dependency_code": ctx.get("dependency_code"), "flow_code": ctx.get("flow_code"),
                    "format_code": ctx.get("format_code"),
                    "mapped_fields": ["code", "description", "category", "mandatory", "source_reference"],
                    "method": "DETERMINISTIC_CONFIGURED_DOCUMENT_STRUCTURE",
                },
                "confidence": "STRUCTURED_SOURCE", "status": "REVIEW_REQUIRED",
            })
        return out

    def _extract_tables(self, blocks: list[dict[str, Any]], cfg: dict[str, Any], **ctx: Any) -> list[dict[str, Any]]:
        """Map explicitly configured table columns. No row is interpreted unless its
        header/column contract is present in configuration."""
        rules = cfg.get("table_rules") or []
        if isinstance(rules, dict): rules = [rules]
        out: list[dict[str, Any]] = []
        groups: dict[tuple[Any, Any], list[dict[str, Any]]] = {}
        for block in blocks:
            if block.get("table") is None or block.get("row") is None or block.get("cell") is None:
                continue
            groups.setdefault((block.get("page"), block.get("table")), []).append(block)
        for rule in rules:
            if not isinstance(rule, dict): continue
            id_headers = {self._norm(x) for x in rule.get("id_headers", [])}
            desc_headers = {self._norm(x) for x in rule.get("description_headers", [])}
            mandatory_headers = {self._norm(x) for x in rule.get("mandatory_headers", [])}
            for _, cells in groups.items():
                rows: dict[Any, list[dict[str, Any]]] = {}
                for b in cells: rows.setdefault(b.get("row"), []).append(b)
                header_map: dict[Any, str] = {}
                for row_no in sorted(rows):
                    vals = sorted(rows[row_no], key=lambda x: str(x.get("cell")))
                    normalized = [self._norm(str(x.get("text") or "")) for x in vals]
                    if any(v in id_headers for v in normalized) and any(v in desc_headers for v in normalized):
                        for b, v in zip(vals, normalized):
                            if v in id_headers: header_map[self._column_key(b.get("cell"))] = "id"
                            elif v in desc_headers: header_map[self._column_key(b.get("cell"))] = "description"
                            elif v in mandatory_headers: header_map[self._column_key(b.get("cell"))] = "mandatory"
                        continue
                    if not header_map: continue
                    mapped: dict[str, dict[str, Any]] = {}
                    for b in vals:
                        role = header_map.get(self._column_key(b.get("cell")))
                        if role: mapped[role] = b
                    if "description" not in mapped: continue
                    description = str(mapped["description"].get("text") or "").strip()
                    if len(description) < 8: continue
                    official_id = self._explicit_identifier(str(mapped.get("id", {}).get("text") or "")) if mapped.get("id") else None
                    raw_id = str(mapped.get("id", {}).get("text") or "").strip() if mapped.get("id") else ""
                    official_id = official_id or (raw_id.upper() if raw_id and len(raw_id) <= 120 else None)
                    ordinal = len(out) + 1
                    code = official_id or f"SRC-{ctx['source_hash'][:10].upper()}-T{ordinal:03d}"
                    ref_block = mapped["description"]
                    ref = self._location(ref_block, ctx, None, description)
                    ref["table"] = ref_block.get("table"); ref["row"] = ref_block.get("row"); ref["cell"] = ref_block.get("cell"); ref["field"] = "table_row"
                    mandatory = None
                    if mapped.get("mandatory"):
                        mandatory = self._parse_boolean(mapped["mandatory"].get("text"))
                    out.append({
                        "code": code, "official_identifier": official_id, "description": description,
                        "category": str(rule.get("category") or "PRESENTACION").upper(), "mandatory": mandatory,
                        "source_reference": ref,
                        "mapping": {"target": "TenderRequirement", "dependency_code": ctx.get("dependency_code"), "flow_code": ctx.get("flow_code"), "format_code": ctx.get("format_code"), "mapped_fields": ["id", "description", "mandatory", "source_reference"], "method": "DETERMINISTIC_TABLE_COLUMN_MAPPING"},
                        "confidence": "STRUCTURED_TABLE", "status": "REVIEW_REQUIRED",
                    })
        return out

    def _dedupe_candidates(self, candidates: list[dict[str, Any]]) -> list[dict[str, Any]]:
        seen: set[tuple[str, str]] = set(); out = []
        for c in candidates:
            key = (str(c.get("code")), str((c.get("source_reference") or {}).get("text_original")))
            if key in seen: continue
            seen.add(key); out.append(c)
        return out

    @staticmethod
    def _norm(value: str) -> str:
        import unicodedata
        value = unicodedata.normalize("NFKD", value).encode("ascii", "ignore").decode("ascii")
        return re.sub(r"\s+", " ", value.strip().lower())

    @staticmethod
    def _column_key(cell: Any) -> str:
        if isinstance(cell, str):
            m = re.match(r"([A-Z]+)", cell.upper())
            return m.group(1) if m else cell
        return str(cell)

    @staticmethod
    def _parse_boolean(value: Any) -> bool | None:
        v = TenderSourceExtractor._norm(str(value or ""))
        if v in {"si", "yes", "true", "x", "obligatorio", "requerido", "mandatory"}: return True
        if v in {"no", "false", "opcional", "optional"}: return False
        return None

    def _select_format(self, role_cfg: dict[str, Any], format_code: str | None, physical_format: str | None) -> dict[str, Any]:
        formats = role_cfg.get("formats") or {}
        selected = str(format_code or "").upper()
        physical = str(physical_format or "").upper()
        for key in (selected, physical, "*"):
            if key and isinstance(formats.get(key), dict):
                return formats[key]
        return role_cfg

    def _explicit_identifier(self, text: str) -> str | None:
        value = text.strip()
        # Preserve a complete explicit identifier, including when it appears inside
        # a sentence such as “Se sustituye R-02…”.
        for rx in (
            re.compile(r"\b((?:REQUISITO|REQ|R|FORMATO|ANEXO|AP[ÉE]NDICE|APENDICE)\s*[-.:#]?\s*[A-Z0-9][A-Z0-9._/-]{0,30})\b", re.I),
            re.compile(r"\b([A-Z]{1,6}[-_]\d{1,8})\b", re.I),
        ):
            m = rx.search(value)
            if m:
                token = re.sub(r"\s+", "", m.group(1)).strip(" .:;-)").upper()
                # Do not treat a bare prose word such as R as an identifier.
                if len(token) > 1 and (any(ch.isdigit() for ch in token) or token.startswith(("REQ", "REQUISITO", "FORMATO", "ANEXO", "APENDICE"))):
                    return token
        m = re.match(r"^\s*(\d+(?:\.\d+){0,5})[.)]?\s+", value)
        return m.group(1) if m else None

    def _explicit_mandatory(self, block: dict[str, Any], cfg: dict[str, Any]) -> bool | None:
        field = cfg.get("mandatory_field")
        if field and field in block:
            value = str(block.get(field)).strip().upper()
            if value in {"SI", "SÍ", "YES", "TRUE", "OBLIGATORIO", "REQUERIDO", "MANDATORY"}: return True
            if value in {"NO", "FALSE", "OPCIONAL", "OPTIONAL"}: return False
        return None

    def _location(self, block: dict[str, Any], ctx: dict[str, Any], section: str | None, original: str) -> dict[str, Any]:
        ref = {
            "source_id": ctx["source_id"], "source_hash": ctx["source_hash"],
            "source_role": ctx["source_role"], "document": ctx.get("filename") or block.get("document"),
            "document_version": ctx.get("document_version"), "revision": ctx.get("source_revision"),
            "page": block.get("page"), "section": block.get("section") or section,
            "table": block.get("table"), "row": block.get("row"), "cell": block.get("cell"),
            "field": block.get("field"), "line_start": block.get("line"), "line_end": block.get("line"),
            "text_original": original, "authority": "TENDER_SOURCE",
            "confirmation_required": True, "extraction_method": "DETERMINISTIC_CONFIGURED_DOCUMENT_STRUCTURE",
        }
        return {k: v for k, v in ref.items() if v is not None}

    def _text_blocks(self, text: str) -> list[dict[str, Any]]:
        return [{"text": line, "line": n} for n, line in enumerate(text.splitlines(), start=1)]

    def _pdf(self, filename: str, media_type: str, content: bytes) -> ExtractedSource:
        import pdfplumber
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(content))
        blocks: list[dict[str, Any]] = []
        chunks: list[str] = []
        scanned_pages: list[int] = []
        with pdfplumber.open(io.BytesIO(content)) as pdf:
            for page_no, page in enumerate(pdf.pages, start=1):
                text = page.extract_text(x_tolerance=2, y_tolerance=3) or ""
                tables = page.extract_tables() or []
                if not text.strip():
                    scanned_pages.append(page_no)
                chunks.append(text)
                for line_idx, raw in enumerate(text.splitlines(), start=1):
                    blocks.append({"text": raw, "page": page_no, "line": sum(len(x.splitlines()) for x in chunks[:-1]) + line_idx, "field": "pdf_text"})
                for table_no, table in enumerate(tables, start=1):
                    for row_no, row in enumerate(table, start=1):
                        for col_no, value in enumerate(row, start=1):
                            if value is None or not str(value).strip(): continue
                            blocks.append({"text": str(value).strip(), "page": page_no, "table": table_no, "row": row_no, "cell": col_no, "field": "pdf_table_cell"})
        # OCR is a real fallback only for pages for which the PDF contains no text layer.
        if scanned_pages:
            blocks.extend(self._ocr_pdf_pages(content, scanned_pages))
            chunks.extend([f"[OCR PAGE {p}]" for p in scanned_pages])
            for b in blocks:
                if b.get("field") == "pdf_ocr_text":
                    chunks.append(b["text"])
        text = "\n\n".join(chunks)
        return ExtractedSource(filename, media_type, text[: self.MAX_TEXT_CHARS], len(reader.pages), {
            "format": "pdf", "blocks": blocks, "scanned_pages": scanned_pages,
            "ocr_used": bool(scanned_pages), "tables_extracted": sum(1 for b in blocks if b.get("field") == "pdf_table_cell"),
        })

    def _ocr_pdf_pages(self, content: bytes, pages: list[int]) -> list[dict[str, Any]]:
        from pdf2image import convert_from_bytes
        import pytesseract
        images = convert_from_bytes(content, dpi=250, first_page=min(pages), last_page=max(pages))
        out: list[dict[str, Any]] = []
        for page_no, image in zip(range(min(pages), max(pages) + 1), images):
            if page_no not in pages: continue
            try:
                langs = set(pytesseract.get_languages(config=""))
            except Exception:
                langs = {"eng"}
            lang = "spa+eng" if {"spa", "eng"}.issubset(langs) else ("spa" if "spa" in langs else "eng")
            data = pytesseract.image_to_data(image, lang=lang, output_type=pytesseract.Output.DICT, config="--psm 6")
            grouped: dict[tuple[int, int, int], list[tuple[int, str]]] = {}
            n = len(data.get("text", []))
            for i in range(n):
                txt = str(data["text"][i] or "").strip()
                if not txt: continue
                key = (int(data["block_num"][i]), int(data["par_num"][i]), int(data["line_num"][i]))
                grouped.setdefault(key, []).append((int(data["left"][i]), txt))
            for line_no, words in enumerate(grouped.values(), start=1):
                text = " ".join(w for _, w in sorted(words))
                out.append({"text": text, "page": page_no, "line": line_no, "field": "pdf_ocr_text", "ocr": True})
        return out

    def _docx(self, filename: str, media_type: str, content: bytes) -> ExtractedSource:
        from docx import Document
        document = Document(io.BytesIO(content))
        parts, blocks = [], []
        line_no = 0
        for paragraph in document.paragraphs:
            if paragraph.text:
                line_no += 1; parts.append(paragraph.text)
                blocks.append({"text": paragraph.text, "line": line_no, "field": "docx_paragraph"})
        for table_no, table in enumerate(document.tables, start=1):
            for row_no, row in enumerate(table.rows, start=1):
                for cell_no, cell in enumerate(row.cells, start=1):
                    text = cell.text.strip()
                    if not text: continue
                    line_no += 1; parts.append(text)
                    blocks.append({"text": text, "line": line_no, "table": table_no, "row": row_no, "cell": cell_no, "field": "docx_table_cell"})
        return ExtractedSource(filename, media_type, "\n".join(parts)[: self.MAX_TEXT_CHARS], None, {"format": "docx", "blocks": blocks, "tables": len(document.tables)})

    def _xlsx(self, filename: str, media_type: str, content: bytes) -> ExtractedSource:
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        parts, blocks = [], []
        line_no = 0
        for ws in wb.worksheets:
            line_no += 1; marker = f"[SHEET:{ws.title}]"; parts.append(marker)
            blocks.append({"text": marker, "line": line_no, "field": "sheet", "table": ws.title})
            for row_no, row in enumerate(ws.iter_rows(), start=1):
                for cell in row:
                    value = "" if cell.value is None else str(cell.value)
                    if not value.strip(): continue
                    line_no += 1; parts.append(value)
                    blocks.append({"text": value, "line": line_no, "table": ws.title, "row": row_no, "cell": cell.coordinate, "field": "xlsx_cell"})
        return ExtractedSource(filename, media_type, "\n".join(parts)[: self.MAX_TEXT_CHARS], None, {"format": "xlsx", "sheets": wb.sheetnames, "blocks": blocks})
