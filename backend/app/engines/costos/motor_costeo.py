# Copyright © 2026 Cristian Rodriguez
# All rights reserved.
# Unauthorized copying, modification, distribution, or use is prohibited
# without prior written permission.

"""
Motor de costeo y presupuestos programables.
Portado de Megalodon v3.1 (megalodon_costos_v3_1.py, clases ConceptoAPU /
MotorPreciosBIM) hacia una arquitectura de dataclasses con Decimal.
"""
from dataclasses import dataclass, field
from typing import Dict, List, Optional, Any
from decimal import Decimal, ROUND_HALF_UP
import math

from app.core.constants import CONSTANTES
from app.core.errors import MegalodonException, ErrorCode
from app.engines.costos.parametros import ParametrosCosteoSnapshot

TWO_PLACES = Decimal("0.01")


def _money(value: Decimal) -> Decimal:
    return value.quantize(TWO_PLACES, rounding=ROUND_HALF_UP)


@dataclass
class InsumoCosteo:
    """Insumo para cálculo de costos (material, mano de obra, equipo...)."""
    clave: str
    descripcion: str
    tipo: str  # MATERIAL, MANO_OBRA, EQUIPO, SUBCONTRATO, HERRAMIENTA
    unidad: str
    cantidad: Decimal
    precio_unitario: Decimal
    rendimiento: Decimal = Decimal("1.0")

    @property
    def importe(self) -> Decimal:
        return _money(self.cantidad * self.precio_unitario)


@dataclass
class ConceptoCosteo:
    """Concepto APU (Análisis de Precios Unitarios).

    `cantidad` es la relación de este concepto dentro de la partida que lo
    contiene (p.ej. 1 m2 de muro puede requerir 1.0 de "block" + 0.02 de
    "acero de refuerzo" como conceptos separados). Para el caso simple de
    un solo concepto por partida, cantidad=1.0 (default).
    """
    clave: str
    descripcion: str
    unidad: str
    cantidad: Decimal = Decimal("1.0")
    insumos: List[InsumoCosteo] = field(default_factory=list)

    @property
    def costo_directo_unitario(self) -> Decimal:
        if not self.insumos:
            return Decimal("0.00")
        return _money(sum((i.importe for i in self.insumos), Decimal("0.00")))

    @property
    def costo_directo(self) -> Decimal:
        """Alias retrocompatible: costo directo unitario de este concepto."""
        return self.costo_directo_unitario


@dataclass
class PartidaCosteo:
    """Partida de presupuesto para costeo.

    precio_unitario_manual: cuando la partida viene de un tabulador
    gubernamental (p.ej. el tabulador CDMX importado) el precio unitario
    YA es un dato final del catálogo oficial — no debe recalcularse desde
    conceptos/insumos. En ese caso se fija este campo y la propiedad
    `precio_unitario` lo respeta tal cual, sin tocar indirectos/utilidad
    a nivel de concepto (esos factores se aplican una sola vez, a nivel
    de todo el presupuesto, para evitar doble aplicación de indirectos).
    """
    numero: int
    descripcion: str
    unidad: str
    cantidad: Decimal
    conceptos: List[ConceptoCosteo] = field(default_factory=list)
    precio_unitario_manual: Optional[Decimal] = None

    @property
    def precio_unitario(self) -> Decimal:
        if self.precio_unitario_manual is not None:
            return _money(self.precio_unitario_manual)
        if not self.conceptos:
            return Decimal("0.00")
        # BUG ORIGINAL: solo usaba self.conceptos[0].costo_directo, así que
        # cualquier partida con más de un concepto (ej. "muro" = block +
        # acero + aplanado como conceptos separados) perdía silenciosamente
        # el costo de todos los conceptos menos el primero.
        total = sum(
            (c.costo_directo_unitario * c.cantidad for c in self.conceptos),
            Decimal("0.00"),
        )
        return _money(total)

    @property
    def importe(self) -> Decimal:
        return _money(self.cantidad * self.precio_unitario)


@dataclass
class PresupuestoCosteo:
    """Presupuesto completo para costeo."""
    identificador: str
    nombre: str
    parametros: ParametrosCosteoSnapshot
    partidas: List[PartidaCosteo] = field(default_factory=list)

    @property
    def factor_indirecto(self) -> Decimal:
        return self.parametros.factor_indirecto

    @property
    def factor_utilidad(self) -> Decimal:
        return self.parametros.factor_utilidad

    @property
    def factor_impuesto(self) -> Decimal:
        return self.parametros.factor_impuesto

    @property
    def factor_riesgo(self) -> Decimal:
        return self.parametros.factor_riesgo

    @property
    def monto_directo(self) -> Decimal:
        if not self.partidas:
            return Decimal("0.00")
        return _money(sum((p.importe for p in self.partidas), Decimal("0.00")))

    @property
    def monto_indirecto(self) -> Decimal:
        return _money(self.monto_directo * self.factor_indirecto)

    @property
    def subtotal(self) -> Decimal:
        return _money(self.monto_directo + self.monto_indirecto)

    @property
    def monto_utilidad(self) -> Decimal:
        return _money(self.subtotal * self.factor_utilidad)

    @property
    def monto_riesgo(self) -> Decimal:
        return _money(self.subtotal * self.factor_riesgo)

    @property
    def base_impuesto(self) -> Decimal:
        return _money(self.subtotal + self.monto_utilidad + self.monto_riesgo)

    @property
    def monto_impuesto(self) -> Decimal:
        return _money(self.base_impuesto * self.factor_impuesto)

    @property
    def monto_total(self) -> Decimal:
        return _money(self.base_impuesto + self.monto_impuesto)

    def to_dict(self) -> Dict[str, Any]:
        return {
            "identificador": self.identificador,
            "nombre": self.nombre,
            "monto_directo": float(self.monto_directo),
            "monto_indirecto": float(self.monto_indirecto),
            "monto_utilidad": float(self.monto_utilidad),
            "monto_riesgo": float(self.monto_riesgo),
            "monto_impuesto": float(self.monto_impuesto),
            "monto_total": float(self.monto_total),
            "factor_indirecto": float(self.factor_indirecto),
            "factor_utilidad": float(self.factor_utilidad),
            "factor_impuesto": float(self.factor_impuesto),
            "factor_riesgo": float(self.factor_riesgo),
            "parametros_costeo": self.parametros.to_dict(),
            "partidas": [
                {
                    "numero": p.numero,
                    "descripcion": p.descripcion,
                    "unidad": p.unidad,
                    "cantidad": float(p.cantidad),
                    "precio_unitario": float(p.precio_unitario),
                    "importe": float(p.importe),
                }
                for p in self.partidas
            ],
        }


class MotorCosteo:
    """Motor principal de costeo."""

    def __init__(self):
        self.constantes = CONSTANTES

    def calcular_presupuesto(self, presupuesto: PresupuestoCosteo) -> PresupuestoCosteo:
        """Fuerza la evaluación de todas las propiedades calculadas.

        Las propiedades ya son la fuente de verdad (se recalculan solas
        cada vez que se leen), así que esto solo sirve para "tocar" el
        objeto y detectar errores de datos temprano.
        """
        _ = presupuesto.monto_total
        return presupuesto

    def validar_sobrecostos(
        self,
        presupuesto_actual: PresupuestoCosteo,
        presupuesto_base: PresupuestoCosteo,
    ) -> Dict[str, Any]:
        """Valida sobrecostos vs presupuesto base."""
        sobrecosto = presupuesto_actual.monto_total - presupuesto_base.monto_total
        factor = (
            sobrecosto / presupuesto_base.monto_total
            if presupuesto_base.monto_total
            else Decimal("0")
        )

        return {
            "presupuesto_base": float(presupuesto_base.monto_total),
            "presupuesto_actual": float(presupuesto_actual.monto_total),
            "sobrecosto": float(sobrecosto),
            "factor_sobrecosto": float(factor),
            "dentro_tolerancia": factor <= Decimal(str(self.constantes.TOLERANCIA_FACTOR_SOBRECOSTO)),
            "tolerancia_permitida": self.constantes.TOLERANCIA_FACTOR_SOBRECOSTO,
        }

    def generar_excel(self, presupuesto: PresupuestoCosteo) -> bytes:
        """Genera archivo Excel del presupuesto."""
        import io
        from openpyxl import Workbook
        from openpyxl.styles import Font, Alignment, PatternFill

        wb = Workbook()
        ws = wb.active
        ws.title = "Presupuesto"

        headers = ["No.", "Descripción", "Unidad", "Cantidad", "P.U.", "Importe"]
        ws.append(headers)

        header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
        header_font = Font(color="FFFFFF", bold=True)

        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")

        for p in presupuesto.partidas:
            ws.append([
                p.numero,
                p.descripcion,
                p.unidad,
                float(p.cantidad),
                float(p.precio_unitario),
                float(p.importe),
            ])

        ws.append(["", "", "", "", "SUBTOTAL:", float(presupuesto.monto_directo)])
        ws.append(["", "", "", "", "INDIRECTO:", float(presupuesto.monto_indirecto)])
        ws.append(["", "", "", "", "UTILIDAD:", float(presupuesto.monto_utilidad)])
        if presupuesto.factor_riesgo:
            ws.append(["", "", "", "", "RIESGO:", float(presupuesto.monto_riesgo)])
        ws.append(["", "", "", "", "IVA:", float(presupuesto.monto_impuesto)])
        ws.append(["", "", "", "", "TOTAL:", float(presupuesto.monto_total)])

        ws.column_dimensions["A"].width = 8
        ws.column_dimensions["B"].width = 50
        ws.column_dimensions["C"].width = 12
        ws.column_dimensions["D"].width = 15
        ws.column_dimensions["E"].width = 15
        ws.column_dimensions["F"].width = 18

        buffer = io.BytesIO()
        wb.save(buffer)
        return buffer.getvalue()

    def generar_pdf(self, presupuesto: PresupuestoCosteo) -> bytes:
        """Genera PDF del presupuesto -- mismo contenido que generar_excel()
        (tabla de partidas + resumen de totales), con reportlab.

        Antes esto no existía: app/workers/pdf_tasks.py era un stub que
        devolvía pdf_url="" con status SUCCESS falso. reportlab ya estaba
        en pyproject.toml (declarado para PAdES-LT de firma electrónica)
        pero nunca se usaba en ningún lado del proyecto -- confirmado con
        grep antes de escribir esto."""
        import io as _io
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import letter
        from reportlab.lib.units import cm
        from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

        buffer = _io.BytesIO()
        doc = SimpleDocTemplate(
            buffer, pagesize=letter,
            topMargin=1.5 * cm, bottomMargin=1.5 * cm,
            leftMargin=1.5 * cm, rightMargin=1.5 * cm,
        )
        styles = getSampleStyleSheet()
        titulo_style = ParagraphStyle("TituloPresupuesto", parent=styles["Heading1"], fontSize=14, spaceAfter=2)
        subtitulo_style = ParagraphStyle("Subtitulo", parent=styles["Normal"], fontSize=9, textColor=colors.HexColor("#555555"))
        celda_desc_style = ParagraphStyle("CeldaDesc", parent=styles["Normal"], fontSize=8, leading=10)

        elementos = []
        elementos.append(Paragraph(presupuesto.nombre, titulo_style))
        elementos.append(Paragraph(f"Identificador: {presupuesto.identificador}", subtitulo_style))
        elementos.append(Spacer(1, 0.5 * cm))

        header = ["No.", "Descripción", "Unidad", "Cantidad", "P.U.", "Importe"]
        filas = [header]
        for p in presupuesto.partidas:
            filas.append([
                str(p.numero),
                Paragraph(p.descripcion, celda_desc_style),
                p.unidad,
                f"{p.cantidad:,.2f}",
                f"${p.precio_unitario:,.2f}",
                f"${p.importe:,.2f}",
            ])

        tabla = Table(filas, colWidths=[1.2 * cm, 8 * cm, 2 * cm, 2.3 * cm, 2.5 * cm, 2.7 * cm], repeatRows=1)
        tabla.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F4E78")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 8),
            ("ALIGN", (0, 0), (0, -1), "CENTER"),
            ("ALIGN", (3, 0), (-1, -1), "RIGHT"),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#CCCCCC")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F5F7FA")]),
            ("TOPPADDING", (0, 0), (-1, -1), 4),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ]))
        elementos.append(tabla)
        elementos.append(Spacer(1, 0.6 * cm))

        filas_totales = [
            ["Monto directo:", f"${presupuesto.monto_directo:,.2f}"],
            [f"Indirecto ({presupuesto.factor_indirecto:.1%}):", f"${presupuesto.monto_indirecto:,.2f}"],
            [f"Utilidad ({presupuesto.factor_utilidad:.1%}):", f"${presupuesto.monto_utilidad:,.2f}"],
        ]
        if presupuesto.factor_riesgo:
            filas_totales.append([f"Riesgo ({presupuesto.factor_riesgo:.1%}):", f"${presupuesto.monto_riesgo:,.2f}"])
        filas_totales.append([f"IVA ({presupuesto.factor_impuesto:.1%}):", f"${presupuesto.monto_impuesto:,.2f}"])
        filas_totales.append(["TOTAL:", f"${presupuesto.monto_total:,.2f}"])

        tabla_totales = Table(filas_totales, colWidths=[5 * cm, 3.5 * cm], hAlign="RIGHT")
        tabla_totales.setStyle(TableStyle([
            ("FONTSIZE", (0, 0), (-1, -1), 9),
            ("ALIGN", (0, 0), (-1, -1), "RIGHT"),
            ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
            ("FONTSIZE", (0, -1), (-1, -1), 11),
            ("LINEABOVE", (0, -1), (-1, -1), 1, colors.black),
            ("TOPPADDING", (0, -1), (-1, -1), 6),
        ]))
        elementos.append(tabla_totales)

        doc.build(elementos)
        return buffer.getvalue()
